import hashlib
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from consolidate import REQUIRED_COLUMNS, consolidate
from generate_sample_data import build_months

ROOT = Path(__file__).resolve().parents[1]


def row(invoice="INV-202601-0001", date=datetime(2026, 1, 5), cust="000123", name="Sample Co",
        postcode="0800", product="P-0010", qty=2, price=10.0):
    return [invoice, date, cust, name, postcode, product, qty, price]


def write_xlsx(path: Path, rows, columns=REQUIRED_COLUMNS):
    pd.DataFrame(rows, columns=columns).to_excel(path, index=False)
    return path


@pytest.fixture
def dirs(tmp_path):
    inp, out = tmp_path / "input", tmp_path / "output"
    inp.mkdir()
    return inp, out


def test_preserves_leading_zeros_in_output_file(dirs):
    inp, out = dirs
    write_xlsx(inp / "jan.xlsx", [row()])
    result = consolidate(inp, out)
    ws = load_workbook(result.report_path)["Consolidated"]
    values = {h.value: c.value for h, c in zip(ws[1], ws[2])}
    assert values["customer_id"] == "000123"
    assert values["postcode"] == "0800"


def test_restores_ids_stored_as_numbers_with_warning(dirs):
    inp, out = dirs
    write_xlsx(inp / "jan.xlsx", [row(cust=123, postcode=800)])
    result = consolidate(inp, out)
    assert result.consolidated.loc[0, "customer_id"] == "000123"
    assert result.consolidated.loc[0, "postcode"] == "0800"
    assert (result.exceptions["severity"] == "warning").sum() == 2


def test_detects_exact_and_conflicting_duplicates(dirs):
    inp, out = dirs
    write_xlsx(inp / "2026-01.xlsx", [row("INV-202601-0001"), row("INV-202601-0002")])
    write_xlsx(inp / "2026-02.xlsx", [row("INV-202601-0001"), row("INV-202601-0002", qty=9)])
    result = consolidate(inp, out)
    assert len(result.consolidated) == 2
    assert sorted(result.duplicates["duplicate_type"]) == ["Conflicting values", "Exact duplicate"]
    assert result.consolidated.set_index("invoice_id").at["INV-202601-0002", "quantity"] == 2  # first kept


def test_rejects_file_missing_required_column(dirs):
    inp, out = dirs
    write_xlsx(inp / "jan.xlsx", [row()])
    cols = [c for c in REQUIRED_COLUMNS if c != "quantity"]
    write_xlsx(inp / "feb.xlsx", [row()[:6] + row()[7:]], columns=cols)
    result = consolidate(inp, out)
    assert result.rejected_files == ["feb.xlsx"]
    assert result.stats["Files processed"] == 1


def test_invalid_values_are_excluded_and_reported(dirs):
    inp, out = dirs
    write_xlsx(inp / "jan.xlsx", [
        row("INV-202601-0001"),
        row("INV-202601-0002", date="31/02/2026"),
        row("INV-202601-0003", qty="two"),
        row("INV-202601-0004", price=-5),
        row("INV-202601-0005", name=None),
        row("BAD-ID"),
    ])
    result = consolidate(inp, out)
    assert list(result.consolidated["invoice_id"]) == ["INV-202601-0001"]
    errors = result.exceptions[result.exceptions["severity"] == "error"]
    assert set(errors["column"]) == {"invoice_date", "quantity", "unit_price", "customer_name", "invoice_id"}
    assert result.stats["Rows with errors (excluded)"] == 5


def test_normalises_messy_headers_and_text_dates(dirs):
    inp, out = dirs
    messy = ["Invoice ID", " Invoice Date ", "CUSTOMER ID", "Customer Name", "Postcode",
             "Product Code", "Quantity", "Unit Price"]
    write_xlsx(inp / "jan.xlsx", [row(date="05/01/2026")], columns=messy)
    result = consolidate(inp, out)
    assert result.consolidated.loc[0, "invoice_date"] == pd.Timestamp(2026, 1, 5)


def test_input_files_are_never_modified(dirs):
    inp, out = dirs
    path = write_xlsx(inp / "jan.xlsx", [row(), row("INV-202601-0001")])
    before = (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns)
    consolidate(inp, out)
    after = (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns)
    assert before == after
    assert sorted(p.name for p in inp.iterdir()) == ["jan.xlsx"]


@pytest.mark.parametrize("sub", [".", "reports"])
def test_refuses_output_inside_input_folder(dirs, sub):
    inp, _ = dirs
    write_xlsx(inp / "jan.xlsx", [row()])
    with pytest.raises(ValueError):
        consolidate(inp, inp / sub)


def test_sample_data_is_synthetic():
    for df in build_months().values():
        names = df.filter(like="ustomer", axis=1).filter(regex="(?i)name").iloc[:, 0].dropna()
        assert names.str.startswith("Sample ").all()


def test_sample_data_end_to_end(tmp_path):
    result = consolidate(ROOT / "sample_data" / "input", tmp_path)
    s = result.stats
    assert s["Files found"] == 7 and s["Files rejected"] == 1
    assert s["Rows read"] - s["Rows with errors (excluded)"] - s["Duplicate rows removed"] == s["Rows in final report"]
