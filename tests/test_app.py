import hashlib
import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from demo_service import SAMPLE_DIR, describe_inputs, run_sample

ROOT = Path(__file__).resolve().parents[1]


def _hashes():
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in SAMPLE_DIR.glob("*.xlsx")}


def test_report_bytes_are_a_valid_workbook_matching_stats():
    result, report = run_sample()
    wb = load_workbook(BytesIO(report))
    assert wb.sheetnames == ["Summary", "Consolidated", "Duplicates", "Exceptions"]
    assert wb["Consolidated"].max_row - 1 == result.stats["Rows in final report"] == len(result.consolidated)
    assert wb["Duplicates"].max_row - 1 == result.stats["Duplicate rows removed"]
    assert wb["Exceptions"].max_row - 1 == len(result.exceptions)


def test_run_leaves_no_files_and_inputs_unchanged(tmp_path, monkeypatch):
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    before = _hashes()
    run_sample()
    assert list(tmp_path.iterdir()) == []
    assert _hashes() == before


def test_input_summary_reconciles_with_run():
    table = describe_inputs()
    result, _ = run_sample()
    assert len(table) == result.stats["Files found"]
    rejected_rows = table[table["File"].isin(result.rejected_files)]["Rows"].sum()
    assert table["Rows"].sum() - rejected_rows == result.stats["Rows read"]


def test_app_runs_and_shows_computed_results():
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    assert not at.exception
    assert len(at.metric) == 0  # nothing shown before the button is pressed
    at.button[0].click().run()
    assert not at.exception
    stats = run_sample()[0].stats
    shown = {m.label: m.value for m in at.metric}
    assert shown["Rows in final report"] == str(stats["Rows in final report"])
    assert shown["Duplicates removed"] == str(stats["Duplicate rows removed"])
    assert shown["Files processed"] == f"{stats['Files processed']} of {stats['Files found']}"
