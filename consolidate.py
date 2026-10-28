"""Consolidate monthly Excel sales files into one validated, formatted report.

Usage:
    python consolidate.py --input sample_data/input --output sample_output
"""
from __future__ import annotations

import argparse
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

REQUIRED_COLUMNS = [
    "invoice_id", "invoice_date", "customer_id", "customer_name",
    "postcode", "product_code", "quantity", "unit_price",
]
# Fixed-width numeric identifiers: stored as text so leading zeros survive.
ID_WIDTHS = {"customer_id": 6, "postcode": 4}
TEXT_COLUMNS = ["invoice_id", "customer_id", "customer_name", "postcode", "product_code"]
INVOICE_PATTERN = r"INV-\d{6}-\d{4}"
DATE_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y")
REPORT_NAME = "consolidated_report.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13, color="1F4E78")


@dataclass
class Result:
    consolidated: pd.DataFrame
    duplicates: pd.DataFrame
    exceptions: pd.DataFrame
    stats: dict
    report_path: Path
    rejected_files: list = field(default_factory=list)


def normalise_header(name) -> str:
    return re.sub(r"\s+", "_", str(name).strip().lower())


def parse_date(value):
    if not isinstance(value, str):
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def check_paths(input_dir: Path, output_dir: Path) -> None:
    input_dir, output_dir = input_dir.resolve(), output_dir.resolve()
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Input folder not found: {input_dir}")
    if output_dir == input_dir or input_dir in output_dir.parents:
        raise ValueError("Output folder must be outside the input folder so source files are never overwritten.")


def _exception(source_file, source_row, column, issue, value, severity):
    return {
        "severity": severity, "source_file": source_file, "source_row": source_row,
        "column": column, "issue": issue, "value": "" if pd.isna(value) else str(value),
    }


def read_file(path: Path, exceptions: list):
    """Read one workbook as text. Returns None if required columns are missing."""
    df = pd.read_excel(path, dtype=str, engine="openpyxl")
    df.columns = [normalise_header(c) for c in df.columns]
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        exceptions.append(_exception(path.name, None, ", ".join(missing),
                                     "Missing required column(s); file rejected", "", "error"))
        return None
    df = df[REQUIRED_COLUMNS].copy()
    df.insert(0, "source_file", path.name)
    df.insert(1, "source_row", df.index + 2)  # Excel row number (header is row 1)
    return df.dropna(how="all", subset=REQUIRED_COLUMNS)


def validate_rows(df: pd.DataFrame, exceptions: list) -> pd.DataFrame:
    """Log problems to `exceptions`; return only error-free rows, typed."""
    df = df.copy()
    errors = pd.Series(False, index=df.index)

    def flag(mask, column, issue, severity="error"):
        nonlocal errors
        for i in df.index[mask]:
            exceptions.append(_exception(df.at[i, "source_file"], int(df.at[i, "source_row"]),
                                         column, issue, df.at[i, column], severity))
        if severity == "error":
            errors |= mask

    present = {}
    for col in REQUIRED_COLUMNS:
        df[col] = df[col].str.strip()
        missing = df[col].isna() | (df[col] == "")
        flag(missing, col, "Missing value")
        present[col] = ~missing

    for col, width in ID_WIDTHS.items():
        cleaned = df[col].str.replace(r"\.0$", "", regex=True)  # numeric cells may read as "123.0"
        numeric = cleaned.str.fullmatch(r"\d+", na=False)
        short = numeric & (cleaned.str.len() < width)
        flag(short, col, f"Leading zeros missing; restored to {width} digits", "warning")
        df[col] = cleaned.where(~short, cleaned.str.zfill(width))
        flag(present[col] & ~df[col].str.fullmatch(rf"\d{{{width}}}", na=False), col, f"Must be {width} digits")

    flag(present["invoice_id"] & ~df["invoice_id"].str.fullmatch(INVOICE_PATTERN, na=False),
         "invoice_id", "Invalid format (expected INV-YYYYMM-NNNN)")

    dates = df["invoice_date"].map(parse_date)
    flag(present["invoice_date"] & dates.isna(), "invoice_date", "Invalid date")

    qty = pd.to_numeric(df["quantity"], errors="coerce")
    flag(present["quantity"] & ~((qty > 0) & (qty % 1 == 0)), "quantity", "Must be a positive whole number")

    price = pd.to_numeric(df["unit_price"], errors="coerce")
    flag(present["unit_price"] & ~(price >= 0), "unit_price", "Must be a non-negative number")

    valid = df[~errors].copy()
    valid["invoice_date"] = pd.to_datetime(dates[~errors])
    valid["quantity"] = qty[~errors].astype(int)
    valid["unit_price"] = price[~errors].astype(float)
    valid["line_total"] = (valid["quantity"] * valid["unit_price"]).round(2)
    return valid


def split_duplicates(df: pd.DataFrame, exceptions: list):
    """Keep the first occurrence of each invoice_id; return (unique, duplicates)."""
    id_dup = df.duplicated("invoice_id", keep="first")
    exact_dup = df.duplicated(REQUIRED_COLUMNS, keep="first")
    first = df[~id_dup].set_index("invoice_id")
    dups = df[id_dup].copy()
    dups.insert(0, "duplicate_type", ["Exact duplicate" if e else "Conflicting values"
                                      for e in exact_dup[id_dup]])
    dups.insert(1, "first_seen", [f"{first.at[i, 'source_file']} row {first.at[i, 'source_row']}"
                                  for i in dups["invoice_id"]])
    for _, row in dups[dups["duplicate_type"] == "Conflicting values"].iterrows():
        exceptions.append(_exception(row["source_file"], int(row["source_row"]), "invoice_id",
                                     f"Same invoice_id as {row['first_seen']} with different values; "
                                     "first kept, review required", row["invoice_id"], "warning"))
    return df[~id_dup].reset_index(drop=True), dups.reset_index(drop=True)


def consolidate(input_dir, output_dir, show_progress: bool = False) -> Result:
    start = time.perf_counter()
    input_dir, output_dir = Path(input_dir), Path(output_dir)
    check_paths(input_dir, output_dir)
    files = sorted(p for p in input_dir.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"No .xlsx files found in {input_dir}")

    exceptions, frames, rejected, rows_read = [], [], [], 0
    for path in tqdm(files, desc="Processing files", unit="file", disable=not show_progress):
        df = read_file(path, exceptions)
        if df is None:
            rejected.append(path.name)
            continue
        rows_read += len(df)
        frames.append(validate_rows(df, exceptions))

    valid = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=REQUIRED_COLUMNS)
    consolidated, duplicates = split_duplicates(valid, exceptions)
    consolidated = consolidated.sort_values(["invoice_date", "invoice_id"]).reset_index(drop=True)
    exc_df = pd.DataFrame(exceptions, columns=["severity", "source_file", "source_row", "column", "issue", "value"])

    error_rows = exc_df[(exc_df["severity"] == "error") & exc_df["source_row"].notna()]
    stats = {
        "Files found": len(files),
        "Files processed": len(files) - len(rejected),
        "Files rejected": len(rejected),
        "Rows read": rows_read,
        "Rows with errors (excluded)": error_rows[["source_file", "source_row"]].drop_duplicates().shape[0],
        "Duplicate rows removed": len(duplicates),
        "Rows in final report": len(consolidated),
        "Warnings": int((exc_df["severity"] == "warning").sum()),
        "Total revenue (AUD)": round(float(consolidated["line_total"].sum()), 2) if len(consolidated) else 0.0,
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / REPORT_NAME
    stats["Processing time (s)"] = round(time.perf_counter() - start, 2)
    write_report(report_path, consolidated, duplicates, exc_df, stats)
    return Result(consolidated, duplicates, exc_df, stats, report_path, rejected)


def write_report(path: Path, consolidated, duplicates, exceptions, stats) -> None:
    stats_df = pd.DataFrame(list(stats.items()), columns=["Metric", "Value"])
    by_month = (consolidated.assign(month=consolidated["invoice_date"].dt.strftime("%Y-%m"))
                .groupby("month", as_index=False)
                .agg(invoices=("invoice_id", "count"), revenue_aud=("line_total", "sum")))
    by_product = (consolidated.groupby("product_code", as_index=False)
                  .agg(units=("quantity", "sum"), revenue_aud=("line_total", "sum"))
                  .sort_values("revenue_aud", ascending=False))

    sections = [("Processing statistics", stats_df), ("Revenue by month", by_month),
                ("Revenue by product", by_product)]
    header_rows, row = [], 1
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for title, table in sections:
            pd.DataFrame([[title]]).to_excel(writer, sheet_name="Summary", startrow=row - 1, header=False, index=False)
            table.to_excel(writer, sheet_name="Summary", startrow=row, index=False)
            header_rows.append((row, row + 1))  # (title row, header row), 1-based
            row += len(table) + 3
        consolidated.to_excel(writer, sheet_name="Consolidated", index=False)
        duplicates.to_excel(writer, sheet_name="Duplicates", index=False)
        exceptions.to_excel(writer, sheet_name="Exceptions", index=False)

    wb = load_workbook(path)
    summary = wb["Summary"]
    for title_row, head_row in header_rows:
        summary.cell(title_row, 1).font = TITLE_FONT
        for cell in summary[head_row]:
            if cell.value is not None:
                cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for r in summary.iter_rows():
        for cell in r:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    for name in ("Consolidated", "Duplicates", "Exceptions"):
        ws = wb[name]
        for cell in ws[1]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        headers = [c.value for c in ws[1]]
        for idx, header in enumerate(headers, start=1):
            fmt = {"invoice_date": "yyyy-mm-dd", "unit_price": "#,##0.00", "line_total": "#,##0.00"}.get(header)
            if header in TEXT_COLUMNS:
                fmt = "@"
            if fmt:
                for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    cell.number_format = fmt
    for ws in wb.worksheets:
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)
    wb.save(path)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--input", required=True, type=Path, help="Folder containing monthly .xlsx files")
    parser.add_argument("--output", required=True, type=Path, help="Folder for the report (outside input)")
    args = parser.parse_args(argv)
    try:
        result = consolidate(args.input, args.output, show_progress=True)
    except (FileNotFoundError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print("\nProcessing statistics")
    for key, value in result.stats.items():
        print(f"  {key:<30} {value:,}" if isinstance(value, (int, float)) else f"  {key:<30} {value}")
    if result.rejected_files:
        print(f"  Rejected: {', '.join(result.rejected_files)}")
    print(f"\nReport saved: {result.report_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
